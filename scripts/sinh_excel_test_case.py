# -*- coding: utf-8 -*-
import pickle, collections, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GOC = pathlib.Path(__file__).resolve().parent.parent
NODES = GOC / 'scripts' / '_nodes.txt'
hang = pickle.load(open(GOC / 'scripts' / '_rows.pkl', 'rb'))

# Số ca chạy thật của từng hàm (tính cả parametrize)
dem = collections.Counter()
for dong in open(NODES):
    p = dong.strip().split('::')
    if len(p) >= 3:
        ten = p[2].split('[')[0]
        dem[(p[0].split('/')[-1], p[1], ten)] += 1
for h in hang:
    h['Số ca chạy'] = dem[(h['Tệp'], h['Nhóm chức năng'], h['Tên test case'])] or 1

FONT = 'Arial'
XANH = PatternFill('solid', fgColor='1F4E79')
NHAT = PatternFill('solid', fgColor='DDEBF7')
VIEN = Border(*[Side(style='thin', color='BFBFBF')] * 4)
MAU_TANG = {'Unit': 'E2EFDA', 'Integration': 'FFF2CC', 'E2E': 'FCE4D6'}

wb = Workbook()

# ============ Sheet 1: Danh mục Test Case ============
ws = wb.active
ws.title = 'Danh muc Test Case'
COT = ['STT', 'Tầng', 'Tệp', 'Nhóm chức năng', 'Tên test case', 'Số ca chạy',
       'Ý nghĩa và cách thực thi']
RONG = [6, 13, 26, 32, 46, 10, 110]

ws.append(COT)
for i, c in enumerate(ws[1], 1):
    c.font = Font(name=FONT, bold=True, size=11, color='FFFFFF')
    c.fill = XANH
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = VIEN
    ws.column_dimensions[get_column_letter(i)].width = RONG[i-1]
ws.row_dimensions[1].height = 30

for h in hang:
    ws.append([h['STT'], h['Tầng'], h['Tệp'], h['Nhóm chức năng'],
               h['Tên test case'], h['Số ca chạy'], h['Ý nghĩa và cách thực thi']])
    r = ws.max_row
    for j in range(1, 8):
        o = ws.cell(row=r, column=j)
        o.font = Font(name=FONT, size=10)
        o.border = VIEN
        o.alignment = Alignment(vertical='top', wrap_text=(j == 7))
    ws.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='top')
    ws.cell(row=r, column=6).alignment = Alignment(horizontal='center', vertical='top')
    ws.cell(row=r, column=2).fill = PatternFill('solid', fgColor=MAU_TANG[h['Tầng']])
    ws.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='top')
    ws.cell(row=r, column=5).font = Font(name='Consolas', size=10)
    ws.cell(row=r, column=7).font = Font(name=FONT, size=9)

ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:G{ws.max_row}'
CUOI = ws.max_row

# ============ Sheet 2: Tổng hợp ============
s2 = wb.create_sheet('Tong hop')
s2.append(['BẢNG TỔNG HỢP KIỂM THỬ'])
s2['A1'].font = Font(name=FONT, bold=True, size=16, color='1F4E79')
s2.merge_cells('A1:D1')
s2.append([])
s2.append(['Tầng kiểm thử', 'Số hàm test', 'Số ca chạy', 'Tỉ lệ'])
for c in s2[3]:
    c.font = Font(name=FONT, bold=True, size=11, color='FFFFFF')
    c.fill = XANH
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = VIEN

D = f"'Danh muc Test Case'"
for i, tang in enumerate(['Unit', 'Integration', 'E2E'], start=4):
    s2.cell(row=i, column=1, value=tang)
    s2.cell(row=i, column=2, value=f'=COUNTIF({D}!$B$2:$B${CUOI},A{i})')
    s2.cell(row=i, column=3, value=f'=SUMIF({D}!$B$2:$B${CUOI},A{i},{D}!$F$2:$F${CUOI})')
    s2.cell(row=i, column=4, value=f'=C{i}/$C$7')
    s2.cell(row=i, column=4).number_format = '0.0%'
    s2.cell(row=i, column=1).fill = PatternFill('solid', fgColor=MAU_TANG[tang])
s2.cell(row=7, column=1, value='TỔNG CỘNG')
s2.cell(row=7, column=2, value='=SUM(B4:B6)')
s2.cell(row=7, column=3, value='=SUM(C4:C6)')
s2.cell(row=7, column=4, value='=C7/C7')
s2.cell(row=7, column=4).number_format = '0.0%'
for r in range(4, 8):
    for j in range(1, 5):
        o = s2.cell(row=r, column=j)
        o.font = Font(name=FONT, size=11, bold=(r == 7))
        o.border = VIEN
        if j > 1:
            o.alignment = Alignment(horizontal='center')
s2['A7'].fill = NHAT; s2['B7'].fill = NHAT; s2['C7'].fill = NHAT; s2['D7'].fill = NHAT

s2.append([]); s2.append(['Thống kê theo tệp'])
s2['A9'].font = Font(name=FONT, bold=True, size=13, color='1F4E79')
s2.append(['Tệp test', 'Số hàm test', 'Số ca chạy'])
for c in s2[10]:
    c.font = Font(name=FONT, bold=True, size=11, color='FFFFFF')
    c.fill = XANH; c.border = VIEN
    c.alignment = Alignment(horizontal='center')
for i, tep in enumerate(sorted({h['Tệp'] for h in hang}), start=11):
    s2.cell(row=i, column=1, value=tep)
    s2.cell(row=i, column=2, value=f'=COUNTIF({D}!$C$2:$C${CUOI},A{i})')
    s2.cell(row=i, column=3, value=f'=SUMIF({D}!$C$2:$C${CUOI},A{i},{D}!$F$2:$F${CUOI})')
    for j in range(1, 4):
        o = s2.cell(row=i, column=j)
        o.font = Font(name=FONT, size=10); o.border = VIEN
        if j > 1:
            o.alignment = Alignment(horizontal='center')
for col, w in zip('ABCD', [42, 16, 14, 12]):
    s2.column_dimensions[col].width = w

# ============ Sheet 3: Chú giải fixture ============
import importlib.util
spec = importlib.util.spec_from_file_location('g', '/tmp/claude-0/-home-user-DoAn/4bd0e4cd-a226-54ea-a89a-898e70ea1209/scratchpad/gen_xlsx.py')
s3 = wb.create_sheet('Chu giai Fixture')
s3.append(['DỮ LIỆU MẪU (FIXTURE) DÙNG TRONG TEST'])
s3['A1'].font = Font(name=FONT, bold=True, size=14, color='1F4E79')
s3.merge_cells('A1:B1')
s3.append([])
s3.append(['Tên fixture', 'Ý nghĩa'])
for c in s3[3]:
    c.font = Font(name=FONT, bold=True, size=11, color='FFFFFF')
    c.fill = XANH; c.border = VIEN
    c.alignment = Alignment(horizontal='center')

# đọc lại từ điển FIXTURE trong gen_xlsx.py
import re as _re
src = pathlib.Path('/tmp/claude-0/-home-user-DoAn/4bd0e4cd-a226-54ea-a89a-898e70ea1209/scratchpad/gen_xlsx.py').read_text(encoding='utf-8')
block = src.split('FIXTURE = {')[1].split('\n}')[0]
pairs = _re.findall(r"'([^']+)':\s*'([^']*)'", block)
for ten, y in pairs:
    s3.append([ten, y])
    r = s3.max_row
    s3.cell(row=r, column=1).font = Font(name='Consolas', size=10)
    s3.cell(row=r, column=2).font = Font(name=FONT, size=10)
    for j in (1, 2):
        s3.cell(row=r, column=j).border = VIEN
        s3.cell(row=r, column=j).alignment = Alignment(vertical='top', wrap_text=(j == 2))
s3.column_dimensions['A'].width = 26
s3.column_dimensions['B'].width = 62
s3.freeze_panes = 'A4'

out = pathlib.Path('docs/DANH-MUC-TEST-CASE.xlsx')
wb.save(out)
print(f'✓ {out}  |  {len(hang)} hàm test, {sum(h["Số ca chạy"] for h in hang)} ca chạy')
